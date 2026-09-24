"""Create the plain-language explanation requested by the user."""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / '第二题研究策略与模型方案_通俗讲解.docx'
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin, section.bottom_margin = Cm(2.1), Cm(2.0)
section.left_margin, section.right_margin = Cm(2.3), Cm(2.3)
section.header_distance, section.footer_distance = Cm(0.9), Cm(0.9)

def font_style(style, name, size, bold=False):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.insert(0, fonts)
    for key in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
        fonts.set(qn('w:' + key), name)
    for key in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
        fonts.attrib.pop(qn('w:' + key), None)

font_style(doc.styles['Normal'], '宋体', 11)
normal = doc.styles['Normal'].paragraph_format
normal.line_spacing = 1.2
normal.space_after = Pt(6)
normal.widow_control = True
font_style(doc.styles['Title'], '微软雅黑', 21, True)
doc.styles['Title'].paragraph_format.space_after = Pt(9)
doc.styles['Title'].paragraph_format.line_spacing = 1.08
font_style(doc.styles['Subtitle'], '微软雅黑', 10)
doc.styles['Subtitle'].font.italic = False
doc.styles['Subtitle'].paragraph_format.space_after = Pt(15)
font_style(doc.styles['Heading 1'], '微软雅黑', 14, True)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(12)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(7)
doc.styles['Heading 1'].paragraph_format.keep_with_next = True
font_style(doc.styles['Header'], '微软雅黑', 9)
font_style(doc.styles['Footer'], '微软雅黑', 9)

header = section.header.paragraphs[0]
header.text = '脑电图计算模型  第二题研究方案'
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('第 ')
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
footer.add_run(' 页')

def p(text, bold=False):
    para = doc.add_paragraph()
    para.add_run(text).bold = bold
    return para

def bullet(text):
    para = p('•  ' + text)
    para.paragraph_format.left_indent = Cm(0.25)
    para.paragraph_format.first_line_indent = Cm(-0.25)
    return para

def heading(text):
    para = doc.add_heading(text, 1)
    if text.startswith(('三 ', '五 ')):
        para.paragraph_format.page_break_before = True

doc.add_paragraph('第二题研究策略与模型方案\n通俗讲解', 'Title')
doc.add_paragraph('基于第一题 V7 与 Q1 结果', 'Subtitle')
p('整个方案可以理解为：先解释“左右三角为什么可能产生不同脑电”，再检查“这种不同是否真的能帮助我们区分左右”。', True)
p('第一题已经尽量清理了信号。第二题就是在此基础上，把“图形—大脑响应—电极记录”这条关系建立起来。本文说明研究思路与模型含义，尚未给出第二题的参数求解和识别性能。')

heading('一 为什么左右三角会产生不同响应')
p('左右三角都有三条边，单纯统计“有几条边、有哪些倾斜方向”，可能很难区分它们。真正有用的是：尖角在哪里，边与边怎样排列。')
p('可以把模型中的神经元群想象成一些“特征探测器”：')
bullet('一组对“尖角朝左”的排列更敏感；')
bullet('一组对“尖角朝右”的排列更敏感；')
bullet('还有一些对两种三角共有的轮廓都响应。')
p('这与题目给出的人脸例子是一致的：重要的不只是有没有眼睛和嘴巴，还包括它们的位置关系。')
p('因此，我们提出的第一个假设是：左右三角激活的神经元群组合不同，随后产生的脑电也可能不同。', True)
p('这里不能简单说“左三角激活右脑、右三角激活左脑”。三角朝向和它出现在视野哪个位置，是两回事。')

heading('二 怎样把神经元群活动变成脑电曲线')
p('一个群体受到刺激后，响应通常不是瞬间出现、瞬间消失，而是经历一定的上升、变化和衰减。不同群体的响应速度也可能不同。')
p('所以，我们用几种简单的时间曲线表示这些响应：较快的响应、较迟缓的响应，以及持续时间更长的慢响应。然后把它们按不同权重叠加起来，生成模型预测的脑电曲线。')
p('例如，某种刺激可能产生较强的快响应和较弱的慢响应；另一种刺激则可能慢响应更明显。即使使用同一组基本曲线，叠加比例不同，最后的波形也会不同。')
p('这些时间曲线受到群体动力学的约束。它们是对复杂活动的简化，暂时不能直接认定某条慢曲线就来自海马体。')

heading('三 为什么还要考虑三个观测位置')
p('可以把 Fz、F3、F4 三个电极理解成放在不同位置的麦克风：同一组声源，到不同麦克风上的混合比例不同。类似地，脑内的多个活动源经过组织传导，在三个电极上形成不同的混合信号。')
p('我们没有足够的信息把所有脑内来源逐一找出来，但可以先研究三种容易理解的变化。')

table = doc.add_table(rows=1, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
table.columns[0].width, table.columns[1].width = Cm(7.3), Cm(9.1)
rows = [
    ('观察方式', '想了解什么'),
    ('三个通道的共同变化', '它们是否一起上升或下降？'),
    ('Fz 相对 F3、F4 的变化', '中间位置是否与两侧不同？'),
    ('F3 与 F4 的差异', '两侧记录是否存在不对称？'),
]
for index, values in enumerate(rows):
    row = table.rows[0] if index == 0 else table.add_row()
    row.cells[0].width, row.cells[1].width = Cm(7.3), Cm(9.1)
    trpr = row._tr.get_or_add_trPr()
    no_split = OxmlElement('w:cantSplit')
    trpr.append(no_split)
    if index == 0:
        repeat = OxmlElement('w:tblHeader')
        trpr.append(repeat)
    for cell, text in zip(row.cells, values):
        cell.text = text
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        tcpr = cell._tc.get_or_add_tcPr()
        borders = OxmlElement('w:tcBorders')
        for edge in ['top', 'left', 'bottom', 'right']:
            e = OxmlElement('w:' + edge)
            e.set(qn('w:val'), 'single')
            e.set(qn('w:sz'), '5')
            e.set(qn('w:color'), 'D9D9D9')
            borders.append(e)
        tcpr.append(borders)
        margins = OxmlElement('w:tcMar')
        for edge, width in [('top', 110), ('bottom', 110), ('left', 130), ('right', 130)]:
            e = OxmlElement('w:' + edge)
            e.set(qn('w:w'), str(width))
            e.set(qn('w:type'), 'dxa')
            margins.append(e)
        tcpr.append(margins)
        if index == 0:
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), 'E8EEF3')
            tcpr.append(shade)
        for para in cell.paragraphs:
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.15
            for run in para.runs:
                run.font.size = Pt(10.5)
                run.bold = index == 0

after_table = p('这一步很重要，因为 Q1 的图提示：有些左右刺激差异表现为三个通道共同变化，有些表现为中间和两侧不同。只看 F3−F4，可能把有用信息抵消掉。')
after_table.paragraph_format.space_before = Pt(7)
p('这三种观察方式也只是信号分解，不能直接当成三个脑区。')

heading('四 最终模型到底在计算什么')
p('用一句话表示，就是：')
p('某个方向的脑电响应 = 两种三角共有的响应 + 该方向带来的额外变化。', True)
p('对于同一个人、同一个项目，模型同时拟合左右两组曲线：先找出两组都存在的共同部分，再找出随方向发生变化的部分，并观察这些变化分别出现在什么空间模式、什么时间尺度上。')
p('这样，模型不仅给出两条拟合曲线，还能说明：左右差异主要体现为共同响应强度不同，还是中间与两侧的关系不同，或者响应的快慢成分比例不同。这些都是可以用数据检验的解释。')

heading('五 怎样得到区分左右的特征')
p('原始一个试次有很多采样点，直接比较整段曲线比较复杂。我们把它压缩成九个数：')
p('三种空间变化 × 三种时间响应 = 九个特征。', True)
p('例如，其中一个数表示“F3 与 F4 差异中的慢响应有多强”，另一个表示“三个通道共同的快响应有多强”。')
p('每个试次都能转换成这九个数。随后，用一个简单分类器检查：左刺激试次和右刺激试次，在这些数上是否存在稳定区别。')
p('关键是，计算未知试次的九个数时，不需要提前知道它是左还是右。', True)

heading('六 怎样判断这个方案是否可信')
p('最重要的不是模型能不能把已有平均曲线画得很像，而是：用一部分试次建立模型，能不能解释和识别后面没有参与建模的试次？', True)
p('我们会重点做几件事：')
bullet('用前面的试次训练，检查后面的试次，判断差异是否随时间保持。')
bullet('打乱左右标签重新训练，检查真实结果是否明显好于偶然结果。')
bullet('去掉某一种空间或时间成分，看看性能是否下降，判断它是否有用。')
bullet('换项目、换受试者检查，了解模型能用到什么范围。')
p('这里需要特别注意第一题的 V7：它清理信号时已经知道左右方向。因此，V7 很适合帮助我们描述已知条件下的波形；但检验“能否识别未知方向”时，需要重新使用不依赖测试方向的处理流程。')
p('第一问已经发现两组数据的左右差异在前后时段并不稳定。所以第二问完全可能得到这样的结果：平均波形可以解释，但稳定区分左右的证据仍不足。如实识别这一点，也是模型检验的重要成果。')

doc.core_properties.title = '第二题研究策略与模型方案通俗讲解'
doc.core_properties.subject = '基于第一题 V7 与 Q1 结果的脑电图建模方案说明'
doc.core_properties.author = ''
doc.core_properties.keywords = '脑电图, 第二题, V7, 研究策略, 通俗讲解'
for root in [doc.styles.element, doc.element]:
    for borders in list(root.iter(qn('w:pBdr'))):
        borders.getparent().remove(borders)
doc.save(OUTPUT)
print(OUTPUT)
