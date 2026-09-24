"""Render the saved discussion summary as an editable Word document."""
from pathlib import Path
import re
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parent
doc = Document(ROOT / '第二题研究策略与模型方案_通俗讲解.docx')
body = doc.element.body
for item in list(body):
    if item.tag != qn('w:sectPr'):
        body.remove(item)
doc.styles['Normal'].paragraph_format.line_spacing = 1.15
doc.styles['Normal'].paragraph_format.space_after = Pt(5)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(11)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(6)
doc.styles['Heading 2'].font.name = '微软雅黑'
doc.styles['Heading 2'].font.size = Pt(11.5)
doc.styles['Heading 2'].font.bold = True
doc.styles['Heading 2'].font.color.rgb = doc.styles['Heading 1'].font.color.rgb
doc.styles['Heading 2'].element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '微软雅黑')
doc.styles['Heading 2'].paragraph_format.keep_with_next = True
doc.styles['Heading 2'].paragraph_format.space_before = Pt(8)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(5)

def clean(s):
    return s.replace('**', '').replace('`', '')

def paragraph(text):
    p = doc.add_paragraph()
    # Preserve emphasized lead paragraphs while leaving the prose editable.
    for j, part in enumerate(re.split(r'(\*\*.*?\*\*)', text)):
        run = p.add_run(clean(part))
        run.bold = part.startswith('**') and part.endswith('**')
    if text.startswith('- '):
        p.text = '• ' + text[2:]
    return p

def add_table(lines):
    rows = [[clean(x.strip()) for x in line.strip().strip('|').split('|')]
            for line in lines if not re.fullmatch(r'[|\s:\-]+', line)]
    nc = len(rows[0])
    widths = [5.3, 11.1] if nc == 2 else [7.1, 3.1, 3.1, 3.1]
    table = doc.add_table(rows=0, cols=nc)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for col, width in zip(table.columns, widths):
        col.width = Cm(width)
    for idx, values in enumerate(rows):
        row = table.add_row()
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if idx == 0:
            trpr.append(OxmlElement('w:tblHeader'))
        for cell, value, width in zip(row.cells, values, widths):
            cell.width = Cm(width)
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for edge in ['top', 'left', 'bottom', 'right']:
                border = OxmlElement('w:' + edge)
                for key, val in [('val', 'single'), ('sz', '5'), ('color', 'D9D9D9')]:
                    border.set(qn('w:' + key), val)
                borders.append(border)
            props.append(borders)
            mar = OxmlElement('w:tcMar')
            for edge, width_twips in [('top', 90), ('bottom', 90), ('left', 110), ('right', 110)]:
                elem = OxmlElement('w:' + edge)
                elem.set(qn('w:w'), str(width_twips))
                elem.set(qn('w:type'), 'dxa')
                mar.append(elem)
            props.append(mar)
            if idx == 0:
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'E8EEF3')
                props.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing = 1.1
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.size = Pt(10)
                    run.bold = idx == 0
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.line_spacing = 0.3
    spacer.add_run().font.size = Pt(3)

lines = (ROOT / '第二题讨论结论与当前完成状态.md').read_text(encoding='utf-8').splitlines()
i = 0
while i < len(lines):
    line = lines[i].strip()
    if not line:
        i += 1
        continue
    if line.startswith('|'):
        table_lines = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            table_lines.append(lines[i])
            i += 1
        add_table(table_lines)
        continue
    if line.startswith('# '):
        doc.add_paragraph(line[2:], 'Title')
        doc.add_paragraph('基于第一题 V7 与 Q1 结果的讨论记录', 'Subtitle')
    elif line.startswith('## '):
        doc.add_heading(line[3:], 1)
    elif line.startswith('### '):
        doc.add_heading(line[4:], 2)
    else:
        paragraph(line)
    i += 1

for root in [doc.styles.element, doc.element]:
    for borders in list(root.iter(qn('w:pBdr'))):
        borders.getparent().remove(borders)
doc.core_properties.title = '第二题讨论结论与当前完成状态'
doc.core_properties.subject = 'V7 标签使用说明 模型可解释性 特征表示和完成状态'
doc.core_properties.author = ''
out = ROOT / '第二题讨论结论与当前完成状态.docx'
doc.save(out)
print(out)
